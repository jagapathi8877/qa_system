"""Data preservation and bounded work regressions from the independent audit."""
import io,json
import pytest
import ingestion,preprocessing

@pytest.mark.parametrize('encoding',['utf-8-sig','utf-16','utf-32','cp1252'])
def test_text_encodings_preserve_facts(encoding):
    text='Café station director is Zoë. The station records 37.50 units.'
    assert ingestion.extract_text('facts.txt',io.BytesIO(text.encode(encoding)))==text

def test_nested_json_owners_and_root_scalars_survive():
    text=ingestion.extract_text('x.json',io.BytesIO(json.dumps({'north':{'director':'Isha'},'south':{'director':'Dev'}}).encode()))
    assert 'north.director: Isha' in text and 'south.director: Dev' in text
    assert ingestion.extract_text('x.json',io.BytesIO(b'"The coastal station has 48 sensors."'))=='The coastal station has 48 sensors.'

@pytest.mark.parametrize('raw',[b'name,role\nNacre,Isha,discarded\n',b'name,role\n"unfinished,Isha\n',b'name,name\na,b\n',b'name,\na,b\n'])
def test_bad_csv_rejected_without_silent_data_loss(raw):
    with pytest.raises(ingestion.ExtractionError):ingestion.extract_text('x.csv',io.BytesIO(raw))

def test_literal_none_csv_is_not_a_missing_value():
    assert 'decision: None' in ingestion.extract_text('x.csv',io.BytesIO(b'name,decision\nNacre,None\n'))

def test_docx_order_and_duplicate_distinct_cells():
    from docx import Document
    doc=Document();doc.add_paragraph('Before table.');table=doc.add_table(rows=1,cols=2)
    for cell in table.rows[0].cells:cell.text='Same value'
    doc.add_paragraph('After table.');buf=io.BytesIO();doc.save(buf);buf.seek(0)
    text=ingestion.extract_text('x.docx',buf)
    assert text.index('Before')<text.index('Same value')<text.index('After')
    assert text.count('Same value')==2

@pytest.mark.parametrize('a,b',[('rechargeable','not rechargeable'),('class A uses 5 watts','class B uses 9 watts'),('-10','10'),('37.50','37.75')])
def test_meaning_critical_minimal_pairs_remain_distinct(a,b):
    assert preprocessing.normalize(a)!=preprocessing.normalize(b)

def test_long_sentence_bounded_and_tail_retained():
    passages=preprocessing.build_passages('Calibration pressure measurement '*4000+' Final identifier ZX-739.')
    assert len(passages)>1 and max(map(len,passages))<=preprocessing.PARAGRAPH_MAX_CHARS
    assert 'ZX-739' in passages[-1]

def test_corpus_work_limit_is_explicit():
    with pytest.raises(ValueError,match='processing limit'):
        preprocessing.build_passages('x'*(preprocessing.MAX_CORPUS_CHARS+1))
