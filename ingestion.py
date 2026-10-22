"""
ingestion.py — turns any supported upload into ONE raw text string.

Supported formats: .txt / .md, .pdf, .docx, .pptx, .csv, .json
Everything downstream (preprocessing, TF-IDF indexing) only ever sees plain
text, so supporting a new format means adding exactly one private function
here and registering it in extract_text().
"""

import csv
import io
import json
import os
import re
import zipfile

SUPPORTED_EXTENSIONS = (".txt", ".md", ".pdf", ".docx", ".pptx", ".csv", ".json")


class UnsupportedFormatError(ValueError):
    """Uploaded file has an extension we cannot parse."""


class ExtractionError(RuntimeError):
    """Parser understood the format but could not read the file."""


def extract_text(filename, stream) -> str:
    """Dispatch on file extension and return all extractable text.

    Blocks are joined with blank lines ("\n\n") on purpose: the chunker in
    preprocessing.py treats a blank-line-separated block as one candidate
    passage, so one paragraph / slide / table row survives as one unit.
    """
    ext = os.path.splitext(filename)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(
            f"Unsupported file type '{ext or '(none)'}'. "
            f"Supported: {', '.join(SUPPORTED_EXTENSIONS)}"
        )
    raw = stream.read(25 * 1024 * 1024 + 1)
    if len(raw) > 25 * 1024 * 1024:
        raise ExtractionError('The file exceeds the 25 MiB parsing limit.')
    if not raw:
        raise ExtractionError("The uploaded file is empty.")
    if ext in ('.docx', '.pptx'):
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                entries = archive.infolist()
                if len(entries) > 10000 or sum(x.file_size for x in entries) > 100 * 1024 * 1024:
                    raise ExtractionError('This document exceeds the safe expanded-size limit.')
        except zipfile.BadZipFile as exc:
            raise ExtractionError('The document is not a valid Office archive.') from exc
    if ext in (".txt", ".md"):
        return _from_plain_text(raw)
    if ext == ".pdf":
        return _from_pdf(raw)
    if ext == ".docx":
        return _from_docx(raw)
    if ext == ".pptx":
        return _from_pptx(raw)
    if ext == ".csv":
        return _from_csv(raw)
    return _from_json(raw)


# ---------------------------------------------------------------- plain text

def _decode_text(raw: bytes) -> str:
    """Respect Unicode BOMs before the legacy Windows text fallback."""
    for mark, encoding in ((b'\xff\xfe\x00\x00', 'utf-32'), (b'\x00\x00\xfe\xff', 'utf-32'),
                           (b'\xff\xfe', 'utf-16'), (b'\xfe\xff', 'utf-16')):
        if raw.startswith(mark):
            try:
                return raw.decode(encoding)
            except UnicodeError as exc:
                raise ExtractionError('Invalid Unicode text encoding.') from exc
    if b'\x00' in raw:
        raise ExtractionError('This file contains binary data or text without a supported encoding marker.')
    for encoding in ("utf-8-sig", "utf-8"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("cp1252", errors="replace")


def _from_plain_text(raw: bytes) -> str:
    return _decode_text(raw)


# ---------------------------------------------------------------------- PDF

def _from_pdf(raw: bytes) -> str:
    import pdfplumber

    pages = []
    try:
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
    except ExtractionError:
        raise
    except Exception as exc:  # encrypted/corrupt files land here
        raise ExtractionError(
            f"Could not read this PDF ({exc}). Is it password-protected?"
        ) from exc
    return "\n\n".join(p.strip() for p in pages if p and p.strip())


# --------------------------------------------------------------------- DOCX

def _from_docx(raw: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(io.BytesIO(raw))
    except Exception as exc:
        raise ExtractionError(f"Could not read this Word file ({exc}).") from exc

    blocks = []
    for element in doc.element.body.iterchildren():
        if element.tag.endswith('}p'):
            text = Paragraph(element, doc).text.strip()
            if text:
                blocks.append(text)
        elif element.tag.endswith('}tbl'):
            for row in Table(element, doc).rows:
                cells, previous_cell = [], None
                for cell in row.cells:
                    # Skip only the same merged cell, not two distinct equal values.
                    if cell._tc is previous_cell:
                        continue
                    previous_cell = cell._tc
                    if cell.text.strip():
                        cells.append(cell.text.strip())
                if cells:
                    blocks.append(', '.join(cells))
    return "\n\n".join(blocks)


# --------------------------------------------------------------------- PPTX

def _from_pptx(raw: bytes) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def shape_texts(shapes):
        """Depth-first walk of a slide's shapes (handles grouped shapes)."""
        parts = []
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                parts.extend(shape_texts(shape.shapes))
            elif getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(", ".join(cells))
            elif getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
        return parts

    try:
        prs = Presentation(io.BytesIO(raw))
    except Exception as exc:
        raise ExtractionError(f"Could not read this PowerPoint file ({exc}).") from exc

    def _terminate(text: str) -> str:
        # Slide titles ("Technology Stack") arrive without punctuation and
        # would glue onto the body text, polluting sentence-level answer
        # extraction — terminate each shape so titles split off cleanly.
        return text if re.search(r"[.!?;,]$", text) else text + "."

    slides = []
    for slide in prs.slides:
        parts = shape_texts(slide.shapes)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(notes)
        if parts:
            slides.append("\n".join(_terminate(p) for p in parts))
    return "\n\n".join(slides)


# ---------------------------------------------------------------------- CSV

def _from_csv(raw: bytes) -> str:
    text = _decode_text(raw)
    try:
        reader = csv.DictReader(io.StringIO(text), strict=True)
        if not reader.fieldnames or any(not x.strip() for x in reader.fieldnames):
            raise ExtractionError('CSV columns must have nonempty headers.')
        if len(set(x.strip().casefold() for x in reader.fieldnames)) != len(reader.fieldnames):
            raise ExtractionError('CSV headers must be unique.')
        rows = []
        for row_number, row in enumerate(reader, 2):
            if None in row:
                raise ExtractionError(f'CSV row {row_number} has more cells than column headers.')
            cells = []
            for key, value in row.items():
                if key and value is not None:
                    val_str = str(value).strip()
                    if val_str:
                        cells.append(f"{key.strip()}: {val_str}")
            if cells:
                rows.append(". ".join(cells))
        return "\n\n".join(rows)
    except csv.Error as exc:
        raise ExtractionError(f"Could not parse this CSV ({exc}).") from exc


# --------------------------------------------------------------------- JSON

def _walk_json(node, out: list, path: str = '', depth: int = 0):
    """Keep complete parent paths so repeated keys do not lose their owner."""
    if depth > 64:
        raise ExtractionError('JSON nesting exceeds 64 levels.')
    if isinstance(node, dict):
        for key, value in node.items():
            label = f'{path}.{key}' if path else str(key)
            if isinstance(value, str) and value.strip():
                out.append(f"{label}: {value.strip()}")
            elif isinstance(value, bool):
                out.append(f"{label}: {value}")
            elif isinstance(value, (int, float)):
                out.append(f"{label}: {value}")
            elif isinstance(value, list) and all(isinstance(x, (str, int, float, bool)) for x in value):
                items = [str(x).strip() if isinstance(x, str) else str(x) for x in value]
                items = [x for x in items if x]
                if items:
                    out.append(f"{label}: {', '.join(items)}")
            else:
                _walk_json(value, out, label, depth + 1)
    elif isinstance(node, list):
        for index, item in enumerate(node):
            _walk_json(item, out, f'{path}[{index}]' if path else f'item[{index}]', depth + 1)
    elif node is not None and str(node).strip():
        out.append(f'{path}: {str(node).strip()}' if path else str(node).strip())


def _from_json(raw: bytes) -> str:
    try:
        def invalid_constant(value):
            raise ValueError('Non-finite JSON number: ' + value)
        data = json.loads(_decode_text(raw), parse_constant=invalid_constant)
    except (ValueError, RecursionError) as exc:
        raise ExtractionError(f"Invalid JSON file ({exc}).") from exc
    out = []
    _walk_json(data, out)
    return "\n\n".join(out)
